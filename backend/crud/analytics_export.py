"""Excel workbook builders for the analytics tab exports.

Both ingestion and digestion tables share one grouped/subtotaled report
layout (a classic "control break" report): rows are grouped Date → Group1
(supplier for ingestion, receipt for digestion) → Client → Fabric Code,
with a bold subtotal row closing each level and a grand total row at the
end. Row bands alternate per date group for readability.

All builder functions take a `lang` ("en" | "ar") that selects every label
below via `_labels_for(lang)`; unrecognized values fall back to English. Sheet data
itself (names, fabric codes, etc.) is never translated — only report chrome.
"""

import io
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_HEADER_FILL = PatternFill("solid", fgColor="2F4858")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_GRAND_FILL = PatternFill("solid", fgColor="1B2A38")
_GRAND_FONT = Font(color="FFFFFF", bold=True)
_LEVEL_FILLS = [
    PatternFill("solid", fgColor="BDD7EE"),  # level 0: date subtotal
    PatternFill("solid", fgColor="D9E1F2"),  # level 1: supplier/receipt subtotal
    PatternFill("solid", fgColor="E2EFDA"),  # level 2: client subtotal
    PatternFill("solid", fgColor="F2F2F2"),  # level 3: fabric code subtotal
]
_BAND_FILL = PatternFill("solid", fgColor="FAFAFA")
_THIN = Side(style="thin", color="D8D8D8")
_BORDER = Border(top=_THIN, bottom=_THIN, left=_THIN, right=_THIN)

_INVALID_SHEET_CHARS = set('[]:*?/\\')

_GROUP_KEYS = ["date", "group1", "client_name", "fabric_label"]
_DETAIL_KEYS_BASE = ["color_name", "material_name"]

# ── Labels ───────────────────────────────────────────────────────────────────
_LABELS: Dict[str, Dict[str, str]] = {
    "en": {
        "date": "Date",
        "supplier": "Supplier",
        "receipt": "Receipt",
        "client": "Client",
        "fabric_code": "Fabric Code",
        "color": "Color",
        "material": "Material",
        "ingested_by": "Ingested By",
        "remnant": "Remnant",
        "issued_by": "Issued By",
        "weight": "Weight (kg)",
        "length": "Length (m)",
        "rolls": "Rolls",
        "roll_ingestion": "Roll Ingestion",
        "roll_digestion": "Roll Digestion",
        "ingestion_of": "Ingestion — {account}",
        "digestion_of": "Digestion — {account}",
        "subtotal": "Subtotal — {label}",
        "grand_total": "Grand Total",
        "no_data": "No data",
        "unknown_date": "Unknown Date",
        "unknown_supplier": "Unknown Supplier",
        "undyed_material": "Undyed ({material})",
        "unknown_material": "Unknown Material",
        "supplier_receipt": "Supplier Receipt",
        "internal_receipt": "Internal Receipt",
        "external_receipt": "External Receipt",
        "receipt_fallback": "Receipt",
        "sheet_ingestion": "Ingestion",
        "sheet_digestion": "Digestion",
        "sheet_no_data": "No Data",
        "sheet_account": "Account",
        "dash": "—",
    },
    "ar": {
        "date": "التاريخ",
        "supplier": "المورد",
        "receipt": "الإيصال",
        "client": "العميل",
        "fabric_code": "كود القماش",
        "color": "اللون",
        "material": "الخامة",
        "ingested_by": "أُدخلت بواسطة",
        "remnant": "بقايا",
        "issued_by": "أصدرها",
        "weight": "الوزن (كجم)",
        "length": "الطول (م)",
        "rolls": "عدد اللفات",
        "roll_ingestion": "إدخال اللفات",
        "roll_digestion": "صرف اللفات",
        "ingestion_of": "إدخال — {account}",
        "digestion_of": "صرف — {account}",
        "subtotal": "الإجمالي الفرعي — {label}",
        "grand_total": "الإجمالي الكلي",
        "no_data": "لا توجد بيانات",
        "unknown_date": "تاريخ غير معروف",
        "unknown_supplier": "مورد غير معروف",
        "undyed_material": "غير مصبوغ ({material})",
        "unknown_material": "خامة غير معروفة",
        "supplier_receipt": "إيصال مورد",
        "internal_receipt": "إيصال داخلي",
        "external_receipt": "إيصال خارجي",
        "receipt_fallback": "إيصال",
        "sheet_ingestion": "الإدخال",
        "sheet_digestion": "الصرف",
        "sheet_no_data": "لا توجد بيانات",
        "sheet_account": "حساب",
        "dash": "—",
    },
}


def _labels_for(lang: str) -> Dict[str, str]:
    return _LABELS.get(lang, _LABELS["en"])


def _num(v: Any) -> float:
    return round(float(v or 0), 3)


def _fabric_label(row: Dict[str, Any], labels: Dict[str, str]) -> str:
    if row.get("roll_type") == "dyed":
        return row.get("fabric_code") or f"#{row.get('client_fabric_code_id')}"
    material = row.get("material_name") or labels["unknown_material"]
    return labels["undyed_material"].format(material=material)


def prepare_rows(rows: List[Dict[str, Any]], direction: str, lang: str = "en") -> List[Dict[str, Any]]:
    """Annotate raw crud row dicts with the two display fields the sheet
    writer groups on: group1 (supplier or receipt label) and fabric_label."""
    labels = _labels_for(lang)
    receipt_kind_labels = {
        "supplier": labels["supplier_receipt"],
        "internal": labels["internal_receipt"],
        "external": labels["external_receipt"],
    }
    out = []
    for r in rows:
        if direction == "ingested":
            group1 = r.get("counterparty") or labels["unknown_supplier"]
        else:
            kind = receipt_kind_labels.get(r.get("receipt_kind") or "", labels["receipt_fallback"])
            rid = r.get("receipt_id")
            base = f"{kind} #{rid}" if rid is not None else kind
            target = r.get("counterparty")
            group1 = f"{target} — {base}" if target else base
        out.append({**r, "group1": group1, "fabric_label": _fabric_label(r, labels)})
    return out


def _safe_sheet_name(name: str, used: Optional[set] = None, fallback: str = "Account") -> str:
    cleaned = "".join(c for c in (name or fallback) if c not in _INVALID_SHEET_CHARS).strip() or fallback
    cleaned = cleaned[:31]
    if used is None:
        return cleaned
    base = cleaned[:28]
    candidate = cleaned
    n = 2
    while candidate in used:
        candidate = f"{base} {n}"[:31]
        n += 1
    return candidate


def _display(key_name: str, value: Any, labels: Dict[str, str]) -> str:
    if value is None or value == "":
        return labels["unknown_date"] if key_name == "date" else labels["dash"]
    return str(value)


def _sorted_rows(rows: List[Dict[str, Any]], group_keys: List[str]) -> List[Dict[str, Any]]:
    """Stable multi-key sort: innermost group ascending first, primary
    (date) descending last, with unknown dates always sorted to the end."""
    ordered = list(rows)
    for key in reversed(group_keys[1:]):
        ordered = sorted(ordered, key=lambda r, k=key: (r.get(k) is None, r.get(k) or ""))
    known = [r for r in ordered if r.get(group_keys[0]) is not None]
    unknown = [r for r in ordered if r.get(group_keys[0]) is None]
    known = sorted(known, key=lambda r: r[group_keys[0]], reverse=True)
    return known + unknown


def _write_header(ws: Worksheet, row: int, headers: List[str]) -> None:
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _BORDER


def _write_subtotal_row(
    ws: Worksheet, row: int, level: int, label: str, totals: List[float], total_cols: int, weight_col: int, labels: Dict[str, str],
) -> None:
    c_label = ws.cell(row=row, column=level + 1, value=labels["subtotal"].format(label=label))
    c_label.font = Font(bold=True, italic=True)
    ws.cell(row=row, column=weight_col, value=_num(totals[0]))
    ws.cell(row=row, column=weight_col + 1, value=_num(totals[1]) or None)
    ws.cell(row=row, column=weight_col + 2, value=int(totals[2]))
    fill = _LEVEL_FILLS[min(level, len(_LEVEL_FILLS) - 1)]
    for col in range(1, total_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.border = _BORDER
        if col >= weight_col:
            cell.font = Font(bold=True)


def _write_detail_row(
    ws: Worksheet, row: int, r: Dict[str, Any], group_keys: List[str], detail_keys: List[str], banded: bool, labels: Dict[str, str],
) -> None:
    values = [_display(k, r.get(k), labels) for k in group_keys] + [_display(k, r.get(k), labels) for k in detail_keys]
    col = 1
    for v in values:
        cell = ws.cell(row=row, column=col, value=v)
        cell.border = _BORDER
        if banded:
            cell.fill = _BAND_FILL
        col += 1
    weight_col = col
    weight_cell = ws.cell(row=row, column=weight_col, value=_num(r.get("weight")))
    length_val = r.get("length")
    length_cell = ws.cell(row=row, column=weight_col + 1, value=_num(length_val) if length_val else None)
    rolls_cell = ws.cell(row=row, column=weight_col + 2, value=1)
    for cell in (weight_cell, length_cell, rolls_cell):
        cell.border = _BORDER
        if banded:
            cell.fill = _BAND_FILL


def write_table(
    ws: Worksheet,
    start_row: int,
    title: str,
    group1_header: str,
    rows: List[Dict[str, Any]],
    direction: str = "ingested",
    lang: str = "en",
) -> int:
    """Writes one grouped/subtotaled table starting at start_row. Returns
    the next free row (leaves one blank separator row after the table)."""
    labels = _labels_for(lang)
    group_keys = _GROUP_KEYS
    if direction == "ingested":
        extra_headers, extra_keys = [labels["ingested_by"], labels["remnant"]], ["ingested_by", "remnant_flag"]
    else:
        extra_headers, extra_keys = [labels["issued_by"]], ["issued_by"]
    detail_keys = _DETAIL_KEYS_BASE + extra_keys
    n_group = len(group_keys)
    total_cols = n_group + len(detail_keys) + 3
    weight_col = n_group + len(detail_keys) + 1

    row = start_row
    ws.cell(row=row, column=1, value=title).font = Font(bold=True, size=13)
    row += 1

    headers = (
        [labels["date"], group1_header, labels["client"], labels["fabric_code"], labels["color"], labels["material"]]
        + extra_headers
        + [labels["weight"], labels["length"], labels["rolls"]]
    )
    _write_header(ws, row, headers)
    header_row = row
    row += 1

    if not rows:
        ws.cell(row=row, column=1, value=labels["no_data"]).font = Font(italic=True, color="888888")
        return row + 2

    ordered = _sorted_rows(rows, group_keys)

    level_totals = [[0.0, 0.0, 0] for _ in range(n_group)]
    grand_totals = [0.0, 0.0, 0]
    prev_keys: List[Optional[str]] = [None] * n_group
    band = False

    def close_level(level: int) -> None:
        nonlocal row
        label = _display(group_keys[level], prev_keys[level], labels)
        _write_subtotal_row(ws, row, level, label, level_totals[level], total_cols, weight_col, labels)
        row += 1
        level_totals[level] = [0.0, 0.0, 0]

    for r in ordered:
        keys = [r.get(k) for k in group_keys]
        if prev_keys[0] is not None and keys != prev_keys:
            diff_level = next(i for i in range(n_group) if keys[i] != prev_keys[i])
            for lvl in range(n_group - 1, diff_level - 1, -1):
                close_level(lvl)
            if diff_level == 0:
                band = not band

        _write_detail_row(ws, row, r, group_keys, detail_keys, band, labels)
        row += 1

        weight = _num(r.get("weight"))
        length = _num(r.get("length"))
        for i in range(n_group):
            level_totals[i][0] += weight
            level_totals[i][1] += length
            level_totals[i][2] += 1
        grand_totals[0] += weight
        grand_totals[1] += length
        grand_totals[2] += 1
        prev_keys = keys

    for lvl in range(n_group - 1, -1, -1):
        close_level(lvl)

    ws.cell(row=row, column=1, value=labels["grand_total"]).font = _GRAND_FONT
    for col in range(1, total_cols + 1):
        ws.cell(row=row, column=col).fill = _GRAND_FILL
    ws.cell(row=row, column=weight_col, value=_num(grand_totals[0])).font = _GRAND_FONT
    ws.cell(row=row, column=weight_col + 1, value=_num(grand_totals[1]) or None).font = _GRAND_FONT
    ws.cell(row=row, column=weight_col + 2, value=int(grand_totals[2])).font = _GRAND_FONT
    row += 1

    ws.row_dimensions[header_row].height = 20
    return row + 1


def _autosize(ws: Worksheet, n_cols: int, min_width: int = 10, max_width: int = 42) -> None:
    for col in range(1, n_cols + 1):
        letter = get_column_letter(col)
        width = min_width
        for cell in ws[letter]:
            if cell.value is not None:
                width = max(width, len(str(cell.value)) + 2)
        ws.column_dimensions[letter].width = min(width, max_width)


def _apply_direction(ws: Worksheet, lang: str) -> None:
    """Arabic sheets read right-to-left; column/cell content is unaffected."""
    ws.sheet_view.rightToLeft = lang == "ar"


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_roll_flow_workbook(ingested_rows: List[Dict[str, Any]], digested_rows: List[Dict[str, Any]], lang: str = "en") -> bytes:
    labels = _labels_for(lang)
    wb = Workbook()
    ws_in = wb.active
    ws_in.title = labels["sheet_ingestion"]
    _apply_direction(ws_in, lang)
    write_table(ws_in, 1, labels["roll_ingestion"], labels["supplier"], prepare_rows(ingested_rows, "ingested", lang), direction="ingested", lang=lang)
    _autosize(ws_in, 11)

    ws_out = wb.create_sheet(labels["sheet_digestion"])
    _apply_direction(ws_out, lang)
    write_table(ws_out, 1, labels["roll_digestion"], labels["receipt"], prepare_rows(digested_rows, "digested", lang), direction="digested", lang=lang)
    _autosize(ws_out, 10)

    return _to_bytes(wb)


def _write_account_sheet(ws: Worksheet, account_name: str, ingested_rows: List[Dict[str, Any]], digested_rows: List[Dict[str, Any]], lang: str = "en") -> None:
    labels = _labels_for(lang)
    _apply_direction(ws, lang)
    next_row = write_table(
        ws, 1, labels["ingestion_of"].format(account=account_name), labels["supplier"],
        prepare_rows(ingested_rows, "ingested", lang), direction="ingested", lang=lang,
    )
    write_table(
        ws, next_row, labels["digestion_of"].format(account=account_name), labels["receipt"],
        prepare_rows(digested_rows, "digested", lang), direction="digested", lang=lang,
    )
    _autosize(ws, 11)


def build_single_account_workbook(account_name: str, ingested_rows: List[Dict[str, Any]], digested_rows: List[Dict[str, Any]], lang: str = "en") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = _safe_sheet_name(account_name, fallback=_labels_for(lang)["sheet_account"])
    _write_account_sheet(ws, account_name, ingested_rows, digested_rows, lang)
    return _to_bytes(wb)


def build_all_accounts_workbook(accounts: List[Dict[str, Any]], lang: str = "en") -> bytes:
    """accounts: [{"name": ..., "ingested": [...], "digested": [...]}]"""
    labels = _labels_for(lang)
    wb = Workbook()
    wb.remove(wb.active)
    used_names: set = set()
    for acc in accounts:
        name = _safe_sheet_name(acc["name"], used_names, fallback=labels["sheet_account"])
        used_names.add(name)
        ws = wb.create_sheet(name)
        _write_account_sheet(ws, acc["name"], acc["ingested"], acc["digested"], lang)
    if not accounts:
        wb.create_sheet(labels["sheet_no_data"])
    return _to_bytes(wb)
